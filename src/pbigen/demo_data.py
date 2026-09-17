"""Excel sintético de ventas para las fases 0-2. Determinista (semilla fija).

Tres hojas: Ventas (hechos, dos años), Productos y Clientes (dimensiones).
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

PRODUCTOS = [
    (1, "Portátil", "Equipos", 950.0),
    (2, "Monitor", "Periféricos", 240.0),
    (3, "Teclado", "Periféricos", 45.0),
    (4, "Ratón", "Periféricos", 25.0),
    (5, "Docking", "Accesorios", 180.0),
]
CLIENTES = [
    (1, "Acme SA", "Norte", "Empresa"),
    (2, "Bolívar SL", "Sur", "Pyme"),
    (3, "Cantábrico SL", "Norte", "Pyme"),
    (4, "Delta Corp", "Este", "Empresa"),
    (5, "Ébano SA", "Oeste", "Empresa"),
    (6, "Fénix SL", "Centro", "Pyme"),
    (7, "Girasol SL", "Sur", "Pyme"),
    (8, "Horizonte SA", "Centro", "Empresa"),
]


def generate_ventas(path: Path, years: tuple[int, ...] = (2024, 2025), seed: int = 7) -> tuple[int, float]:
    """Escribe Ventas / Productos / Clientes. Devuelve (filas de ventas, importe total)."""
    rng = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Fecha", "ProductoId", "ClienteId", "Unidades", "Importe"])
    total, rows = 0.0, 0
    precios = {pid: precio for pid, _, _, precio in PRODUCTOS}
    for year in years:
        day = date(year, 1, 1)
        growth = 1.0 + 0.08 * (year - years[0])
        while day.year == year:
            for pid, _, _, _ in PRODUCTOS:
                for cid, _, _, _ in CLIENTES:
                    if rng.random() < 0.22:
                        unidades = rng.randint(1, 12)
                        importe = round(unidades * precios[pid] * rng.uniform(0.9, 1.1) * growth, 2)
                        ws.append([day, pid, cid, unidades, importe])
                        total += importe
                        rows += 1
            day += timedelta(days=1)
    for cell in ws["A"][1:]:
        cell.number_format = "yyyy-mm-dd"
    wp = wb.create_sheet("Productos")
    wp.append(["ProductoId", "Producto", "Categoria", "Precio"])
    for row in PRODUCTOS:
        wp.append(list(row))
    wc = wb.create_sheet("Clientes")
    wc.append(["ClienteId", "Cliente", "Region", "Segmento"])
    for row in CLIENTES:
        wc.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return rows, round(total, 2)
